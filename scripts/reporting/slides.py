"""Five editable slides with speaker notes, actual UI evidence and a 28-minute demo plan."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from evidence import ROOT, OUT

NAVY='102C43'; BLUE='176AAD'; TEAL='167D8D'; ORANGE='D77B35'; INK='163B58'; MUTED='526A7C'; PAPER='F4F7FA'; WHITE='FFFFFF'


def box(slide,x,y,w,h,fill,line=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill)
    shape.line.fill.background() if line is None else None
    if line:shape.line.color.rgb=RGBColor.from_string(line)
    return shape


def text(slide,x,y,w,h,value,size=20,color=INK,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y), Inches(w), Inches(h))
    frame=shape.text_frame;frame.word_wrap=True
    frame.margin_left=frame.margin_right=Inches(.03)
    frame.margin_top=frame.margin_bottom=Inches(.02)
    for index,line in enumerate(value.split('\n')):
        paragraph=frame.paragraphs[0] if index==0 else frame.add_paragraph()
        paragraph.text=line;paragraph.font.name='DejaVu Sans';paragraph.font.size=Pt(size)
        paragraph.font.bold=bold;paragraph.font.color.rgb=RGBColor.from_string(color)
        paragraph.space_after=Pt(9)
    return shape


def new(prs,number,eyebrow,title,dark=False):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else PAPER)
    text(slide,.55,.25,11,.35,eyebrow,12, '8FC8DE' if dark else TEAL,True)
    text(slide,.55,.8,12.1,1.2,title,32,WHITE if dark else INK,True)
    text(slide,.55,7.08,11,.22,'SPORTS CENTER  /  SWP391 FA26  /  NHÓM NHANH · BẢO · KHÔI · KHẢI',8,'B3CAD7' if dark else MUTED)
    text(slide,12.1,7.02,.6,.35,f'{number:02d} / 05',10,'B3CAD7' if dark else MUTED)
    return slide


def card(slide,x,y,w,h,kicker,title,body,color=BLUE):
    box(slide,x,y,w,h,WHITE)
    text(slide,x+.2,y+.18,w-.4,.35,kicker,12,color,True)
    text(slide,x+.2,y+.7,w-.4,.8,title,23,INK,True)
    text(slide,x+.2,y+1.65,w-.4,h-1.8,body,17,MUTED)


def build(snapshot):
    prs=Presentation();prs.core_properties.author='Nhóm SWP391';prs.core_properties.title='Sports Center Management System';prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    slide=new(prs,1,'01  GIỚI THIỆU','Hệ thống quản lý\ntrung tâm thể thao',True)
    text(slide,.6,2.35,10.5,.6,'Từ mua gói và đăng ký lớp đến thanh toán, kết quả và báo cáo',21,'C6DDE9')
    box(slide,.6,3.25,11.9,.68,'1B435E')
    text(slide,.8,3.38,11.5,.4,'4 vai trò nghiệp vụ  •  Quyền từ DB  •  Giao dịch có lịch sử  •  Demo local',16,WHITE)
    people=[('Nhanh','Nền tảng và thanh toán\nSprint 1–3'),('Bảo','BE + báo cáo full-stack\nNhận toàn bộ phần Nhanh ở S4'),('Khôi','Hồ sơ full-stack\nLuồng nhân viên và hội viên'),('Khải','FE tổng kết và đánh giá\nKiểm thử, trải nghiệm')]
    for index,(name,role) in enumerate(people):
        x=.65+index*3.07
        text(slide,x,4.55,2.9,.5,name,23,WHITE,True)
        text(slide,x,5.2,2.85,1.05,role,14,'C6DDE9')
    slide.notes_slide.notes_text_frame.text='00:00–01:00. Nhóm trình bày Sports Center Management System cho SWP391 FA26, đề tài 4. Nêu tên bốn thành viên và phạm vi hệ thống. Phần chia việc trên slide là nhóm công việc tiêu biểu; đóng góp đầy đủ xem task sheet và Git. Ở Sprint 4, Bảo nhận toàn bộ phần trước đây giao Nhanh; Khôi làm hồ sơ full-stack, Khải FE tổng kết và kiểm thử.'
    slide=new(prs,2,'02  WHY / CONTEXT','Một thay đổi vận hành chạm nhiều bộ phận')
    for x,kicker,title,body in [(.6,'01  LỊCH VÀ CHỖ','Đăng ký trùng?','Lớp, HLV và phòng có thể xung đột.\nGiữ chỗ phải nhất quán khi nhiều người cùng đăng ký.'),(4.72,'02  TIỀN VÀ QUYỀN LỢI','Tiền về chưa khớp?','Thiếu, thừa, sai mã và chuyển lặp.\nCần biết khi nào được kích hoạt gói hoặc chỗ học.'),(8.84,'03  KẾT QUẢ VÀ LỊCH SỬ','Ai đổi, đổi lúc nào?','Hồ sơ, điểm danh và kết quả cần đúng phạm vi.\nBáo cáo phải giải thích được thu và hoàn.')]:
        card(slide,x,2.1,3.9,3.95,kicker,title,body)
    text(slide,.65,6.28,12,.5,'Bối cảnh nghiệp vụ của đồ án; dữ liệu demo là dữ liệu giả, không phải kết quả khảo sát thực địa.',13,MUTED)
    slide.notes_slide.notes_text_frame.text='01:00–02:00. Dùng một tình huống xuyên suốt: hội viên mua khóa nhưng chuyển thiếu tiền, trong khi chỗ lớp phải được giữ đúng hạn. Nếu chỉ quản lý bằng tin nhắn/bảng tính, trạng thái tiền, chỗ và thông báo có thể lệch nhau. Mục tiêu là kiểm soát các điểm này bằng một luồng có trạng thái, quyền và lịch sử. Không nêu số liệu tiết kiệm thời gian hay doanh thu vì nhóm chưa có đo đạc thực tế.'
    slide=new(prs,3,'03  WHAT + WHO','Một luồng thống nhất, bốn phạm vi trách nhiệm')
    stages=['Chọn gói / khóa','Giữ chỗ + hóa đơn','Thu tiền hợp lệ','Học + tổng kết','Báo cáo + audit']
    for index,label in enumerate(stages):
        x=.6+index*2.5;box(slide,x,2.15,2.25,.72,TEAL if index==2 else BLUE)
        text(slide,x+.08,2.35,2.1,.35,label,13,WHITE,True)
        if index<4:text(slide,x+2.28,2.3,.25,.4,'→',17,MUTED)
    roles=[('HỘI VIÊN','Mua, đăng ký, kết quả\nHồ sơ và đánh giá'),('LỄ TÂN','Tìm và đăng ký hộ\nThu quầy, ghi actor'),('HUẤN LUYỆN VIÊN','Lịch dạy, điểm danh\nHồ sơ, tổng kết'),('QUẢN LÝ','Quyền và đối soát\nKiểm duyệt, báo cáo')]
    for index,(name,body) in enumerate(roles):
        x=.6+index*3.1;box(slide,x,3.45,2.9,1.85,WHITE)
        text(slide,x+.13,3.7,2.64,.5,name,14,TEAL,True);text(slide,x+.13,4.32,2.64,.8,body,14,MUTED)
    text(slide,.65,5.8,12,1,'Gói tự tập ≠ khóa học     |     API kiểm tra quyền và phạm vi     |     Snapshot khi công bố',19,INK,True)
    slide.notes_slide.notes_text_frame.text='02:00–03:30. Giải thích hai sản phẩm độc lập: kỳ gói cho quyền tự tập; khóa/lớp có đăng ký riêng, có thể miễn phí hoặc thu học phí. Lớp thu phí tạo hóa đơn và giữ chỗ; chỉ khi đủ tiền mới BOOKED. Lễ tân làm hộ dùng cùng luật với hội viên nhưng ghi đúng người thao tác. HLV chỉ có học viên theo các buổi thực dạy và đồng ý chia sẻ. Mọi thao tác quan trọng kiểm tra permission ở API, menu không phải ranh giới bảo vệ.'
    slide=new(prs,4,'04  TECHNOLOGY + ARCHITECTURE','Giữ nghiệp vụ nhất quán tại API và dữ liệu',True)
    for x,title,body in [(.6,'TRÌNH DUYỆT','React 19 · Ant Design 6\nRouter · Query · Zustand'),(4.72,'API NGHIỆP VỤ','Node · Express 5 · Zod\nJWT → quyền → service'),(8.84,'DỮ LIỆU','Prisma 7 · MySQL 8\nTransaction · FK · unique')]:
        box(slide,x,2.25,3.9,1.75,'1B435E');text(slide,x+.18,2.52,3.5,.4,title,17,'8FC8DE',True);text(slide,x+.18,3.12,3.5,.65,body,15,WHITE)
    text(slide,4.48,2.95,.25,.4,'→',17,WHITE);text(slide,8.6,2.95,.25,.4,'→',17,WHITE)
    for index,(title,body) in enumerate([('THANH TOÁN','Webhook SePay có xác thực / chống trùng'),('TÁC VỤ NỀN','Hết hạn, danh sách chờ, tổng kết, outbox'),('KẾT QUẢ','Snapshot + revision; xuất PDF / Excel / CSV')]):
        y=4.55+index*.6;text(slide,.7,y,2.5,.4,title,13,'8FC8DE',True);text(slide,3.25,y,9.2,.4,body,16,WHITE)
    slide.notes_slide.notes_text_frame.text='03:30–05:00. Dự án là npm workspaces: core/fe, core/be và core/shared. Browser gửi HTTP JSON; Express xác thực JWT, nạp permission và validate Zod trước service. Service dùng Prisma với MySQL. Tranh chấp chỗ và thanh toán được kiểm tra lại trong transaction có khóa. Worker chạy cùng vòng đời server, chưa phải một dịch vụ queue riêng. SePay gọi webhook vào API; SMTP gửi từ server. Tổng kết công bố có snapshot, revision và audit; báo cáo dùng cùng builder khi xem và xuất file.'
    slide=new(prs,5,'05  LIVE DEMO / 28 PHÚT','Chứng minh luồng chạy được và có bằng chứng')
    steps=[('06–10’','Mua → thiếu tiền → chuyển bù → PAID'),('10–13’','Đăng ký, chặn trùng lịch, xử lý hộ'),('13–16’','Hồ sơ và phạm vi chia sẻ'),('16–21’','Tổng kết → PDF → đánh giá → mở lại'),('21–24’','Báo cáo theo kỳ và xuất Excel')]
    for index,(time,label) in enumerate(steps):
        y=2.15+index*.62;text(slide,.65,y,1.25,.4,time,16,TEAL,True);text(slide,1.95,y,5.6,.45,label,17,INK)
    image=ROOT/'outputs/sprint4-bao/reports-desktop.png'
    if image.exists():
        from PIL import Image
        width,height=Image.open(image).size
        picture=slide.shapes.add_picture(str(image),Inches(7.85),Inches(2.25),width=Inches(4.8),height=Inches(2.8))
        picture.crop_bottom=max(0,1-(2.8/4.8)/(height/width))
    text(slide,7.9,5.35,4.7,.7,'Ảnh giao diện báo cáo Sprint 4\nDữ liệu seed phục vụ demo',12,MUTED)
    box(slide,.65,6.0,11.9,.68,INK);text(slide,.85,6.15,11.5,.4,'24–26’: task + commit + test     •     26–28’: hỏi đáp và dự phòng',16,WHITE)
    slide.notes_slide.notes_text_frame.text='05:00–06:00: giới thiệu thứ tự rồi chuyển trình duyệt. Demo 06:00–24:00 theo runbook. 24:00–26:00 mở một task, commit thật, service và test có liên quan; phân biệt ngày kế hoạch với author/committer date. 26:00–28:00 hỏi đáp. Giữ sẵn profile bốn vai trò, hóa đơn mới và một tổng kết DRAFT. Dùng simulator local đã kiểm tra, không chuyển tiền thật. Nếu webhook không sẵn, đổi sang thu quầy và nói rõ kênh thay thế. Cần cả nhóm diễn tập trực tiếp trước nộp; việc tạo kịch bản không có nghĩa đã diễn tập.'
    prs.save(OUT/'SWP391_ThuyetTrinh_v1.pptx')
