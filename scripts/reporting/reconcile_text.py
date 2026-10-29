"""Correct known superseded rules while preserving historical comparison tables."""
from docx import Document
from evidence import ROOT


def reconcile_rules():
    path=ROOT/'docs/yeu-cau/07_Business_Rules.docx'
    doc=Document(path)
    updates={
      'BR-1.3':'Mỗi hóa đơn gói thanh toán thành công tạo một kỳ Membership riêng. Có kỳ còn giá trị hoặc mua trước thì kỳ mới bắt đầu sau kỳ kết thúc muộn nhất; không cộng đè ngày vào một bản ghi.',
      'BR-1.4':'Khi không còn kỳ gói còn giá trị, kỳ mua mới bắt đầu ở thời điểm thanh toán. Trạng thái và activeUserId được đồng bộ theo ngày; chỉ một kỳ đang hiệu lực của mỗi hội viên.',
      'BR-1.12':'Gói tự tập độc lập với khóa/lớp. Không chặn hoặc cảnh báo quyền đăng ký lớp chỉ vì gói tự tập hết hạn trước ngày kết thúc lớp.',
      'BR-1.16':'Bảo lưu chỉ với kỳ đang hoạt động: mỗi lần ít nhất MEMBERSHIP_FREEZE_MIN_DAYS, tổng ngân sách kỳ theo MEMBERSHIP_FREEZE_MAX_TOTAL_DAYS; chỉ một lịch bảo lưu chưa kết thúc. Cộng đúng ngày thực dừng một lần; không chặn đăng ký lớp chỉ vì bảo lưu gói tự tập.',
      'BR-2.1':'Đăng ký lớp không yêu cầu gói tự tập. Lớp miễn phí đăng ký trực tiếp khi đủ điều kiện; khóa có học phí giữ PENDING_PAYMENT và chỉ BOOKED khi hóa đơn được thanh toán đủ.',
      'BR-2.16':'Danh sách chờ FIFO; chỉ vào chờ khi đủ điều kiện khác và lớp đầy. Lời mời OFFERED còn hạn giữ sức chứa, cần hội viên xác nhận; khóa có học phí vẫn phải thanh toán. Không yêu cầu gói tự tập.',
      'BR-PAY-08':'Giao dịch mới về khi hóa đơn FAILED hoặc PAID đi qua đối soát với lý do LATE hoặc DUPLICATE_TRANSFER; gửi lại cùng ID giao dịch được xử lý idempotent, không cấp quyền lợi hai lần.',
      'BR-PAY-16':'Số đã nhận bằng tổng payment trừ các khoản hoàn áp dụng. Chỉ tự PAID khi tổng đúng số hóa đơn; thiếu giữ PENDING với hạn chuyển bù cấu hình. Quá hạn chuyển phiếu chờ hoàn, không coi tạo lệnh là đã chuyển trả.',
    }
    for table in list(doc.tables):
        if any(c.text.startswith('BR-AI-') for row in table.rows for c in row.cells):
            table._element.getparent().remove(table._element)
            continue
        if any(c.text=='Trước' for c in table.rows[0].cells):
            continue
        for row in table.rows:
            key=row.cells[0].text
            if key in updates and len(row.cells)>=2:
                row.cells[1].text=updates[key]
                if key=='BR-2.1' and len(row.cells)>2:
                    row.cells[2].text='Tách hai sản phẩm: tự tập theo kỳ và khóa học theo đăng ký/thanh toán; tránh áp dụng yêu cầu gói tự tập cũ.'
    for paragraph in list(doc.paragraphs):
        if 'Quy tắc nghiệp vụ cho phần AI' in paragraph.text:
            paragraph._element.getparent().remove(paragraph._element)
    marker='Phạm vi bản triển khai 29/10/2026'
    if not any(p.text.startswith(marker) for p in doc.paragraphs):
        p=doc.paragraphs[0].insert_paragraph_before(marker+': các BR hiện hành đã cập nhật theo mã nguồn. BR-AI và kế hoạch tập luyện là đặc tả ngoài phạm vi đã chốt; không được dùng làm tiêu chí đã triển khai. Bảng Trước/Sau chỉ là lịch sử thay đổi; cột Sau được áp dụng.')
    doc.save(path)


def reconcile_usecases():
    path=ROOT/'docs/yeu-cau/01_Use_Case_Tong_Hop.docx'
    doc=Document(path)
    replacements={
      'A1. Membership hiện tại còn hạn → thời hạn mới được nối tiếp từ ngày hết hạn hiện tại.':'A1. Còn kỳ gói còn giá trị → tạo kỳ mới bắt đầu sau kỳ kết thúc muộn nhất; không cập nhật đè kỳ cũ.',
      'Trong thời gian bảo lưu không đăng ký lớp mới;':'Gói tự tập đang bảo lưu không chặn đăng ký lớp; điều kiện lớp được xét độc lập;',
      'Hội viên có gói đang hiệu lực, không bảo lưu, không bị chặn đăng ký.':'Hội viên ACTIVE, không bị chặn, không trùng lịch, đúng cửa sổ đăng ký và lớp còn sức chứa. Gói tự tập không phải điều kiện. Lớp thu phí dùng luồng mua khóa và giữ chỗ chờ thanh toán.',
      'Lớp đầy: mời vào danh sách chờ (UC-CB-19). Gói hết hạn trước ngày kết thúc lớp: vẫn cho đăng ký nhưng cảnh báo số buổi sau ngày hết hạn và gợi ý gia hạn. Gói đang bảo lưu hoặc hội viên bị chặn: từ chối kèm ngày được đăng ký lại.':'Lớp đầy mà đủ điều kiện khác: vào danh sách chờ (UC-CB-19). Tài khoản bị chặn hoặc trùng lịch: từ chối kèm lý do. Gói tự tập hết hạn/bảo lưu không chặn việc đăng ký lớp. Lớp có học phí chỉ BOOKED sau khi thanh toán đủ.',
      'gói được cộng dồn theo BR-1.3':'một kỳ gói mới được tạo nối tiếp theo BR-1.3',
      'UNDERPAID (đã nhận X, còn thiếu Y)':'PENDING kèm kết quả UNDERPAID (đã nhận X, còn thiếu Y; UNDERPAID không phải enum InvoiceStatus)',
    }
    for table in doc.tables:
        if any(c.text=='Trước' for c in table.rows[0].cells):
            continue
        for row in table.rows:
            for cell in row.cells:
                text=cell.text
                for old,new in replacements.items():
                    text=text.replace(old,new)
                if text!=cell.text:
                    cell.text=text
    marker='Phạm vi hiện hành 29/10/2026'
    if not any(p.text.startswith(marker) for p in doc.paragraphs):
        doc.paragraphs[0].insert_paragraph_before(marker+': không triển khai AI/kế hoạch tập luyện. Các UC tương ứng trong đặc tả gốc là phạm vi đã huỷ; phần đang triển khai gồm tài khoản, gói tự tập, lớp, thanh toán, hồ sơ, tổng kết, đánh giá và báo cáo. Sprint 4 chi tiết tại 08_Sprint4_UseCases_BusinessRules.docx. Bảo nhận toàn bộ task Sprint 4 trước đây giao Nhanh.')
    doc.save(path)
    path=ROOT/'docs/yeu-cau/08_Sprint4_UseCases_BusinessRules.docx'
    doc=Document(path)
    doc.paragraphs[1].text='Phạm vi hiện hành 29/10/2026: không làm AI/kế hoạch tập luyện. Bảo phụ trách toàn bộ nền dữ liệu, BE tổng kết/đánh giá, báo cáo BE/FE và tài liệu/demo; Khôi làm hồ sơ full-stack; Khải làm FE tổng kết/đánh giá và kiểm thử. Task sheet Sprint 4 và bảng phân chia UC đã đồng bộ.'
    doc.save(path)


if __name__=='__main__':
    reconcile_rules()
    reconcile_usecases()
