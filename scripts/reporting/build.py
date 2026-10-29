"""Build local final-submission artifacts from reviewed content and repository evidence."""
import os
import hashlib
import json
import subprocess
from pathlib import Path
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from evidence import ROOT, OUT, collect
import diagrams
import documents
import runbook
import workbook
import slides


def export_pdf():
    targets=[OUT/'SWP391_BaoCaoDoAn_v1.docx',OUT/'SWP391_KichBanDemo_v1.docx',OUT/'SWP391_ThuyetTrinh_v1.pptx']
    subprocess.run([os.environ.get('SWP_LIBREOFFICE','libreoffice'),'-env:UserInstallation=file:///tmp/swp-report-office','--headless','--convert-to','pdf','--outdir',str(OUT),*[str(p) for p in targets]],check=True,timeout=90)
    for path in targets:
        if not path.with_suffix('.pdf').exists():
            raise RuntimeError(f'PDF export missing: {path}')


def validate(snapshot):
    import pymupdf
    deck=Presentation(OUT/'SWP391_ThuyetTrinh_v1.pptx')
    assert len(deck.slides)==5
    assert sum(row[1] for row in __import__('content').DEMO)==28
    assert len(snapshot['models'])==49 and len(snapshot['relations'])==88
    assert all(t['owner']!='Nhanh' and 'be-nhanh' not in str(t['branch']) for t in snapshot['tasks'] if t['id'].startswith('S4'))
    book=load_workbook(OUT/'SWP391_TaskTraceability_v1.xlsx',read_only=True)
    assert book['TaskTraceability'].max_row==len(snapshot['tasks'])+1
    assert book['GitTimeline'].max_row==len(snapshot['commits'])+1
    result={'baseline':snapshot['baseline'],'slides':len(deck.slides),'demoMinutes':28,'activeModels':42,'activeForeignKeys':72,'schemaModels':49,'schemaForeignKeys':88,'tasks':len(snapshot['tasks']),'commits':len(snapshot['commits']),'outputs':{}}
    for name in ['SWP391_BaoCaoDoAn_v1','SWP391_KichBanDemo_v1','SWP391_ThuyetTrinh_v1']:
        pdf=pymupdf.open(OUT/f'{name}.pdf')
        if name.endswith('ThuyetTrinh_v1'):
            assert len(pdf)==5
        assert all(page.get_text().strip() for page in pdf)
        result['outputs'][name]={'pdfPages':len(pdf)}
    json_files=[str(p) for p in OUT.glob('*.json') if p.name!='verification.json']
    subprocess.run(['npx','prettier','--write',*json_files],cwd=ROOT,check=True,capture_output=True)
    for path in OUT.rglob('*'):
        if path.is_file() and path.name!='verification.json':
            result['outputs'][str(path.relative_to(OUT))]=hashlib.sha256(path.read_bytes()).hexdigest()
    (OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in result.items() if k!='outputs'},ensure_ascii=False))


def main():
    snapshot=collect()
    print('Rendering diagrams',flush=True);diagrams.build(snapshot)
    print('Building demo and current flow reference',flush=True);runbook.build(snapshot)
    snapshot=collect()
    print('Building report, workbook and five-slide deck',flush=True)
    documents.build_report(snapshot);workbook.build(snapshot);slides.build(snapshot)
    print('Exporting PDF with LibreOffice',flush=True);export_pdf()
    validate(snapshot)

if __name__=='__main__':
    main()
