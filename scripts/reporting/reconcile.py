"""Apply approved Sprint 4 reassignment and reconcile current documentation."""
import json
import io
import subprocess
import re
from copy import copy
from docx import Document
from evidence import ROOT, OUT, load_sheet


def task_sheet(changes):
    path = ROOT/'docs/ke-hoach/SWP391_TaskSheet_Sprint4.xlsx'
    book = load_sheet(path, data_only=False)
    for sheet in book:
        for row in sheet:
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                before = cell.value
                cell.value = before.replace('sprint-4/be-nhanh','sprint-4/be-bao').replace('Chỉ Nhanh','Chỉ Bảo').replace('nhắn Nhanh','nhắn Bảo').replace('Nhanh / Khôi','Bảo / Khôi')
                cell.value = re.sub(r'Chuyển từ Nhanh sang Bảo theo phân công 29/10/2026\.', 'Bảo phụ trách theo phân công 29/10/2026.', cell.value)
                if cell.value != before:
                    changes.append([str(path.relative_to(ROOT)),sheet.title,cell.coordinate,before,cell.value])
    sheet = book['Task Sprint 4']
    sheet['A2'] = 'Phạm vi hiện hành: không làm AI/kế hoạch tập luyện. Bảo phụ trách toàn bộ phần BE nền dữ liệu, tổng kết, đánh giá, báo cáo BE/FE và tài liệu/demo; Khôi làm hồ sơ full-stack; Khải làm FE tổng kết/đánh giá và E2E. Bắt đầu/Hạn là kế hoạch; ngày Git là bằng chứng thực hiện riêng.'
    sheet['I4'], sheet['J4'] = 'Bắt đầu (dự kiến)', 'Hạn (dự kiến)'
    for row in range(5,sheet.max_row+1):
        if not re.fullmatch(r'S4-\d{2}', str(sheet.cell(row,1).value)):
            continue
        if sheet.cell(row,5).value == 'Nhanh':
            sheet.cell(row,5,'Bảo')
            sheet.cell(row,8,'sprint-4/be-bao')
        if row == 5:
            sheet.cell(row,12,'Nền dữ liệu đã tích hợp main; Bảo phụ trách schema, migration và hợp đồng API. Kiểm tra trên DB tạm và giữ test hồi quy đạt. Ngày trong cột kế hoạch không phải ngày commit thực tế.')
    summary = book['Tổng hợp']
    for row,member,role in [(4,'Bảo','BE + FE báo cáo + nền dữ liệu + tài liệu/demo'),(5,'Khôi','Full-stack hồ sơ + E2E hồ sơ'),(6,'Khải','FE tổng kết/đánh giá + kiểm thử')]:
        summary.cell(row,1,member); summary.cell(row,2,role)
        summary.cell(row,3,f'=COUNTIF(\'Task Sprint 4\'!$E$5:$E$39,A{row})')
        for col,status in [(4,'To Do'),(5,'In Progress'),(6,'Review'),(7,'Done')]:
            summary.cell(row,col,f'=COUNTIFS(\'Task Sprint 4\'!$E$5:$E$39,A{row},\'Task Sprint 4\'!$K$5:$K$39,"{status}")')
        summary.cell(row,8,f'=IFERROR(G{row}/C{row},0)')
    for cell in summary[7]:
        cell.value = None
    for col in range(3,8):
        letter = chr(64+col)
        summary.cell(8,col,f'=SUM({letter}4:{letter}6)')
    summary.cell(8,8,'=IFERROR(G8/C8,0)')
    book.save(path)


def role_split(changes):
    path = ROOT/'docs/ke-hoach/SWP391_UCP_Sprint_Role_Split.xlsx'
    book = load_sheet(path,data_only=False)
    source = load_sheet(ROOT/'docs/ke-hoach/SWP391_TaskSheet_Sprint4.xlsx',data_only=False)['Task Sprint 4']
    tasks = {row[0].value:[c.value for c in row] for row in source if re.fullmatch(r'S4-\d{2}',str(row[0].value))}
    sheet = book['Phân chia UC']
    for row in sheet.iter_rows(min_row=5):
        if row[9].value == 'Sprint 3':
            row[13].value='Review'
        if row[9].value != 'Sprint 4':
            continue
        for cell in row:
            if isinstance(cell.value,str):
                cell.value = cell.value.replace('Nhanh','Bảo').replace('be-nhanh','be-bao')
        if isinstance(row[0].value,int) and 65 <= row[0].value <= 76:
            row[13].value = 'Cancelled'
            row[12].value = 'Ngoài phạm vi Sprint 4 hiện hành từ 26/10/2026; không triển khai AI/kế hoạch tập luyện. Chức năng thay thế #110–126; task chi tiết tại SWP391_TaskSheet_Sprint4.xlsx.'
    groups = [[1,2],[3,5],[4,6],[7,8],[9,10],[11],[12,13],[14,15],[16,17],[18,19],[20,21],[22,23],[24,25],[26,27],[28],[29,30,31],[0,32,33,34]]
    for number,group in enumerate(groups,110):
        selected = [tasks[f'S4-{i:02d}'] for i in group]
        owners = lambda kind: ' / '.join(dict.fromkeys(r[4] for r in selected if r[3]==kind)) or '—'
        row_number = next((r for r in range(5,sheet.max_row+1) if sheet.cell(r,1).value==number),sheet.max_row+1)
        first = selected[0]
        values = [number,('Flow 1 – User & Membership' if number<=114 else 'Flow 4 – Training & Attendance' if number<=119 else 'Flow 3 – Payment & Report' if number<=124 else 'Khác – Tài liệu & chất lượng'),first[1],re.sub(r'^(BE|FE|Kiểm thử|Tài liệu) – ','',first[2]),'Theo permission nghiệp vụ',number,first[5],owners('BE'),owners('FE'),'Sprint 4','03/11 → 16/11 (kế hoạch)','Xem schema.prisma và ERD cuối kỳ','; '.join(f'S4-{i:02d}' for i in group)+'; '+str(first[6]),'In Progress' if 34 in group else 'Review','; '.join(dict.fromkeys(str(r[7]) for r in selected))]
        for column,value in enumerate(values,1):
            target=sheet.cell(row_number,column,value)
            target._style=copy(sheet.cell(5,column)._style)
        sheet.row_dimensions[row_number].height=80
    sheet['A2']='126 hàng chức năng; #65–76 ngoài phạm vi; #110–126 là Sprint 4 hiện hành. Các sprint và lịch sử Git trước đó giữ nguyên.'
    sheet['A3']='Phân công Sprint 4: Bảo nhận toàn bộ phần trước đây giao Nhanh; Khôi hồ sơ full-stack, Khải FE tổng kết/đánh giá và E2E. Trạng thái Review không đồng nghĩa đã duyệt; ngày kế hoạch không thay ngày thực hiện.'
    sheet.cell(4,11,'From–To (kế hoạch)')
    summary=book['Sprint Summary']
    summary['B4']='From–To (kế hoạch)'
    summary['A2']='Tổng hợp phạm vi hiện hành: loại các dòng Cancelled khỏi số chức năng đang thực hiện; task chi tiết và tiến độ Sprint 4 lấy từ task sheet Sprint 4.'
    summary['C8']='Bảo / Khôi (hồ sơ)'; summary['D8']='Khôi / Khải / Bảo (báo cáo)'
    summary['I8']='Bảo: nền dữ liệu, BE tổng kết/đánh giá, báo cáo full-stack, tài liệu/seed/demo. Khôi: hồ sơ full-stack. Khải: FE tổng kết/đánh giá, test. AI/kế hoạch tập #65–76: Cancelled; phạm vi mới #110–126.'
    summary['D15']=str(summary['D15'].value).replace('; AI Sprint 4','; Sprint 4 đã chuyển toàn bộ phần phân công sang Bảo')
    summary['D16']=str(summary['D16'].value)+'; Sprint 4 phụ trách nền dữ liệu, tổng kết, đánh giá, báo cáo full-stack và demo.'
    last=sheet.max_row
    for row in range(5,9):
        summary.cell(row,5,f'=COUNTIFS(\'Phân chia UC\'!J5:J{last},A{row},\'Phân chia UC\'!N5:N{last},"<>Cancelled")')
        for col,level in [(6,'Simple'),(7,'Medium'),(8,'Complex')]:
            summary.cell(row,col,f'=COUNTIFS(\'Phân chia UC\'!J5:J{last},A{row},\'Phân chia UC\'!G5:G{last},"{level}",\'Phân chia UC\'!N5:N{last},"<>Cancelled")')
        for col,status in [(10,'To Do'),(11,'In Progress'),(12,'Review'),(13,'Done')]:
            summary.cell(row,col,f'=COUNTIFS(\'Phân chia UC\'!J5:J{last},A{row},\'Phân chia UC\'!N5:N{last},"{status}")')
    for name in ['Tổng quan đề tài','Sources & Notes']:
        for row in book[name]:
            for cell in row:
                if cell.coordinate in ['D44','D45'] and name=='Tổng quan đề tài':
                    cell.value='Ngoài phạm vi AI/kế hoạch; điểm danh đã có. Sprint 4 hiện hành: hồ sơ, tổng kết, đánh giá, báo cáo.'
        if name=='Sources & Notes':
            book[name]['B14']='Bảo + Khôi (BE hồ sơ) | FE Khôi / Khải / Bảo'
            book[name]['C14']='Kế hoạch 03/11 → 16/11; phân công 29/10: Bảo nhận toàn bộ task trước đây của Nhanh. Task sheet Sprint 4 là nguồn phân công hiện hành; #65–76 Cancelled, #110–126 bổ sung.'
    for row in range(15,19):
        for col,source_col in [(2,'H'),(3,'I')]:
            summary.cell(row,col,f'=COUNTIFS(\'Phân chia UC\'!{source_col}5:{source_col}{last},"*"&A{row}&"*",\'Phân chia UC\'!N5:N{last},"<>Cancelled")')
    for row in range(41,47):
        overview=book['Tổng quan đề tài']
        old=overview.cell(row,5).value
        if isinstance(old,str) and old.startswith('=COUNTIF('):
            prefix,criterion=old[len('=COUNTIF('):-1].rsplit(',',1)
            overview.cell(row,5,f'=COUNTIFS({prefix},{criterion},\'Phân chia UC\'!$N$5:$N${last},"<>Cancelled")')
    for page in book:
        for cells in page:
            for cell in cells:
                if isinstance(cell.value,str) and cell.value.startswith('='):
                    cell.value=cell.value.replace('$113',f'${last}')
    book['Tổng quan đề tài']['B44']='Điểm danh và tổng kết khóa học'
    book['Tổng quan đề tài']['B45']='AI và kế hoạch tập luyện đã loại phạm vi'
    book.save(path)
    changes.append([str(path.relative_to(ROOT)),'Phân chia UC','65–76;110–126','Phạm vi cũ training/AI và phân công Nhanh','Cancelled scope cũ; bổ sung 17 hàng hiện hành, task của Nhanh chuyển Bảo'])


def apply():
    changes=[]
    task_sheet(changes)
    role_split(changes)
    audit()


def audit():
    changes=[]
    for name in ['SWP391_TaskSheet_Sprint4.xlsx','SWP391_UCP_Sprint_Role_Split.xlsx']:
        path='docs/ke-hoach/'+name
        old=load_sheet(io.BytesIO(subprocess.check_output(['git','show','a8ba56c:'+path],cwd=ROOT)),data_only=False)
        new=load_sheet(ROOT/path,data_only=False)
        for page in new:
            prior=old[page.title] if page.title in old.sheetnames else None
            for row in page:
                for cell in row:
                    before=prior[cell.coordinate].value if prior else None
                    if before!=cell.value:
                        changes.append([path,page.title,cell.coordinate,before,cell.value])
    (OUT/'reconciliation.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2,default=str))
    print('Reconciled Sprint 4 assignment and scope:',len(changes))

if __name__=='__main__':
    apply()
